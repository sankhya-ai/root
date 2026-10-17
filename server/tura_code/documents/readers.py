"""Reading the documents the runtime's ``ReadFile`` refuses: PDF, Word, Excel, PowerPoint, CSV.

Each reader returns plain text a model can work with — pages, sheets and
slides marked, tables as rows — plus a short summary for the transcript. A
scanned PDF without a text layer says so rather than returning nothing.
"""

from __future__ import annotations

import csv
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

MAX_CHARS = 60_000
MAX_SHEET_ROWS = 300
MAX_SHEET_COLS = 40

READABLE = (".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".pptx")

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class DocumentError(ValueError):
    pass


def page_numbers(spec: str, count: int) -> list[int]:
    """``"1-3,5"`` → ``[0, 1, 2, 4]`` (zero-based), within ``count`` pages. Empty means all."""
    spec = (spec or "").strip()
    if not spec:
        return list(range(count))
    chosen: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        match = re.fullmatch(r"(\d+)\s*(?:-\s*(\d+|end))?", part)
        if not match:
            raise DocumentError(f"Bad page range {part!r}: use numbers like `2`, `1-3` or `4-end`.")
        first = int(match.group(1))
        last = count if match.group(2) == "end" else int(match.group(2) or first)
        if first < 1 or last > count or first > last:
            raise DocumentError(f"Pages {part} are outside 1-{count}.")
        chosen.extend(n - 1 for n in range(first, last + 1) if n - 1 not in chosen)
    return chosen


def read(path: Path, *, pages: str = "", sheet: str = "", cells: str = "", max_chars: int = MAX_CHARS) -> dict:
    """``{"text", "summary", "meta"}`` for a document."""
    if not path.is_file():
        raise DocumentError(f"No such file: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        result = read_pdf(path, pages)
    elif suffix == ".docx":
        result = read_docx(path)
    elif suffix in (".xlsx", ".xlsm"):
        result = read_xlsx(path, sheet, cells)
    elif suffix in (".csv", ".tsv"):
        result = read_csv(path)
    elif suffix == ".pptx":
        result = read_pptx(path)
    else:
        raise DocumentError(f"read_document reads {', '.join(READABLE)}; use ReadFile for text files.")
    text = result["text"]
    if len(text) > max_chars:
        result["text"] = text[:max_chars] + f"\n\n[truncated: {len(text) - max_chars} more characters; ask for fewer pages or a smaller range]"
        result["meta"]["truncated"] = True
    return result


def read_pdf(path: Path, pages: str = "") -> dict:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError("The PDF is password-protected.")
        count = len(reader.pages)
    except PdfReadError as error:
        raise DocumentError(f"Could not open the PDF: {error}") from error
    chosen = page_numbers(pages, count)
    parts, empty = [], 0
    for index in chosen:
        text = (reader.pages[index].extract_text() or "").strip()
        empty += not text
        parts.append(f"--- page {index + 1} ---\n{text or '[no text on this page]'}")
    meta: dict = {"pages": count, "read": len(chosen)}
    fields = reader.get_fields() or {}
    if fields:
        meta["form_fields"] = {name: str((field or {}).get("/V") or "") for name, field in fields.items()}
        parts.append("--- form fields ---\n" + "\n".join(f"{name}: {value}" for name, value in meta["form_fields"].items()))
    if chosen and empty == len(chosen):
        meta["no_text_layer"] = True
        parts.append("[This PDF has no text layer: it is probably scanned. Text cannot be extracted without OCR.]")
    summary = f"{len(chosen)} of {count} pages" + (f", {len(fields)} form fields" if fields else "")
    return {"text": "\n\n".join(parts), "summary": summary, "meta": meta}


def _xml(archive: zipfile.ZipFile, name: str) -> ElementTree.Element:
    try:
        return ElementTree.fromstring(archive.read(name))
    except KeyError as error:
        raise DocumentError(f"Not a valid document: {name} is missing.") from error
    except ElementTree.ParseError as error:
        raise DocumentError(f"Malformed document: {error}") from error


def _zip(path: Path) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(path)
    except zipfile.BadZipFile as error:
        raise DocumentError(f"Not a valid Office document: {error}") from error


def read_docx(path: Path) -> dict:
    with _zip(path) as archive:
        body = _xml(archive, "word/document.xml").find(f"{_W}body")
    lines: list[str] = []
    paragraphs = tables = 0
    for block in list(body) if body is not None else []:
        if block.tag == f"{_W}p":
            text = "".join(t.text or "" for t in block.iter(f"{_W}t"))
            style = block.find(f"{_W}pPr/{_W}pStyle")
            level = (style.get(f"{_W}val") or "") if style is not None else ""
            heading = re.fullmatch(r"Heading(\d)", level)
            lines.append(("#" * int(heading.group(1)) + " " + text) if heading and text else text)
            paragraphs += 1
        elif block.tag == f"{_W}tbl":
            tables += 1
            for row in block.iter(f"{_W}tr"):
                cells = ["".join(t.text or "" for t in cell.iter(f"{_W}t")).strip() for cell in row.iter(f"{_W}tc")]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return {"text": text or "[the document has no text]", "summary": f"{paragraphs} paragraphs, {tables} tables", "meta": {"paragraphs": paragraphs, "tables": tables}}


def _cell_text(cell, value) -> str:
    raw = cell.value
    if raw is None:
        return ""
    if isinstance(raw, str) and raw.startswith("="):
        return f"{raw} (= {value})" if value is not None else raw
    return str(raw)


def read_xlsx(path: Path, sheet: str = "", cells: str = "") -> dict:
    from openpyxl import load_workbook

    try:
        formulas = load_workbook(str(path), data_only=False, read_only=False)
        values = load_workbook(str(path), data_only=True, read_only=False)
    except Exception as error:  # noqa: BLE001 - openpyxl raises many kinds
        raise DocumentError(f"Could not open the workbook: {error}") from error
    names = formulas.sheetnames
    if sheet and sheet not in names:
        raise DocumentError(f"No sheet {sheet!r}; the sheets are {', '.join(names)}.")
    parts = []
    for name in [sheet] if sheet else names:
        ws, wv = formulas[name], values[name]
        dims = ws.dimensions if ws.max_row and ws.max_column else "empty"
        if cells:
            try:
                grid = ws[cells]
            except ValueError as error:
                raise DocumentError(f"Bad range {cells!r}: {error}") from error
            # One cell, one row or column, or a grid of rows.
            if not isinstance(grid, tuple):
                rows = ((grid,),)
            elif grid and isinstance(grid[0], tuple):
                rows = grid
            else:
                rows = (grid,)
        else:
            rows = ws.iter_rows(min_row=1, max_row=min(ws.max_row, MAX_SHEET_ROWS), max_col=min(ws.max_column, MAX_SHEET_COLS))
        lines = []
        for row in rows:
            texts = [(c.coordinate, _cell_text(c, wv[c.coordinate].value)) for c in row if hasattr(c, "coordinate")]
            filled = [f"{coordinate}={text}" for coordinate, text in texts if text]
            if filled:
                lines.append("  ".join(filled))
        more = "" if cells or ws.max_row <= MAX_SHEET_ROWS else f"\n[{ws.max_row - MAX_SHEET_ROWS} more rows: read a range with `cells`]"
        merged = f"\nmerged: {', '.join(str(r) for r in ws.merged_cells.ranges)}" if ws.merged_cells.ranges else ""
        parts.append(f"--- sheet {name} ({dims}) ---\n" + ("\n".join(lines) or "[empty]") + more + merged)
    summary = f"{len(names)} sheet{'s' if len(names) != 1 else ''}: {', '.join(names)}"
    return {"text": "\n\n".join(parts), "summary": summary, "meta": {"sheets": names}}


def read_csv(path: Path) -> dict:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        rows = list(csv.reader(handle, delimiter=delimiter))
    shown = rows[:MAX_SHEET_ROWS * 2]
    text = "\n".join(f"{n}: " + " | ".join(row) for n, row in enumerate(shown, start=1))
    if len(rows) > len(shown):
        text += f"\n[{len(rows) - len(shown)} more rows]"
    return {"text": text or "[empty]", "summary": f"{len(rows)} rows", "meta": {"rows": len(rows)}}


def read_pptx(path: Path) -> dict:
    with _zip(path) as archive:
        slides = sorted(
            (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)),
        )
        parts = []
        for number, name in enumerate(slides, start=1):
            root = _xml(archive, name)
            texts = []
            for paragraph in root.iter(f"{_A}p"):
                line = "".join(t.text or "" for t in paragraph.iter(f"{_A}t")).strip()
                if line:
                    texts.append(line)
            parts.append(f"--- slide {number} ---\n" + ("\n".join(texts) or "[no text]"))
    return {"text": "\n\n".join(parts) or "[no slides]", "summary": f"{len(slides)} slides", "meta": {"slides": len(slides)}}
