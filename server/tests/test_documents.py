"""Documents: reading PDF/Office/CSV files and editing spreadsheets and PDFs, unit and end to end."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from conftest import of_type, requires_runtime, run_turn, text, tool_call

from tura_code.documents import pdfs, readers, sheets
from tura_code.documents.readers import DocumentError


def make_pdf(path: Path, pages: list[str]) -> Path:
    operations = [{"op": "add_page"} for _ in pages]
    operations += [{"op": "add_text", "page": n, "text": body} for n, body in enumerate(pages, start=1)]
    pdfs.edit(path, operations)
    return path


def make_form(path: Path) -> Path:
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject, TextStringObject

    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    field = DictionaryObject(
        {
            NameObject("/FT"): NameObject("/Tx"),
            NameObject("/T"): TextStringObject("name"),
            NameObject("/V"): TextStringObject(""),
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/F"): NumberObject(4),
            NameObject("/Rect"): ArrayObject([NumberObject(72), NumberObject(700), NumberObject(300), NumberObject(720)]),
            NameObject("/DA"): TextStringObject("/Helv 12 Tf 0 g"),
        }
    )
    ref = writer._add_object(field)
    page[NameObject("/Annots")] = ArrayObject([ref])
    writer._root_object[NameObject("/AcroForm")] = DictionaryObject({NameObject("/Fields"): ArrayObject([ref])})
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def make_docx(path: Path) -> Path:
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    body = (
        f'<w:document {w}><w:body>'
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Release notes</w:t></w:r></w:p>'
        "<w:p><w:r><w:t>Fixed the </w:t></w:r><w:r><w:t>login bug.</w:t></w:r></w:p>"
        "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Area</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Owner</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
        "</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", body)
    return path


def test_pdf_pages_text_and_ranges(tmp_path):
    pdf = make_pdf(tmp_path / "r.pdf", ["Alpha page", "Beta page", "Gamma page"])
    whole = readers.read(pdf)
    assert whole["meta"]["pages"] == 3 and "Beta page" in whole["text"]
    some = readers.read(pdf, pages="2-end")
    assert "Alpha" not in some["text"] and "Gamma page" in some["text"]
    with pytest.raises(DocumentError):
        readers.read(pdf, pages="4")


def test_pdf_edits_merge_delete_reorder_rotate_metadata_encrypt(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", ["one", "two"])
    b = make_pdf(tmp_path / "b.pdf", ["three"])
    result = pdfs.edit(a, [{"op": "merge", "files": [str(b)]}, {"op": "reorder", "order": "3,1,2"}, {"op": "delete_pages", "pages": "3"}, {"op": "metadata", "title": "Merged"}])
    assert result["pages"] == 2
    assert "three" in pdfs.page_text(a, 1) and "one" in pdfs.page_text(a, 2)
    pdfs.edit(a, [{"op": "keep_pages", "pages": "2"}], output=tmp_path / "part.pdf")
    assert "one" in pdfs.page_text(tmp_path / "part.pdf", 1)
    secret = tmp_path / "secret.pdf"
    pdfs.edit(a, [{"op": "encrypt", "password": "pw"}], output=secret)
    with pytest.raises(DocumentError, match="password"):
        readers.read(secret)
    with pytest.raises(DocumentError, match="no pages"):
        pdfs.edit(a, [{"op": "delete_pages", "pages": "1-2"}])


def test_pdf_form_fill(tmp_path):
    form = make_form(tmp_path / "form.pdf")
    assert readers.read(form)["meta"]["form_fields"] == {"name": ""}
    pdfs.edit(form, [{"op": "fill_form", "fields": {"name": "Ada Lovelace"}}])
    assert readers.read(form)["meta"]["form_fields"] == {"name": "Ada Lovelace"}
    with pytest.raises(DocumentError, match="No form fields named"):
        pdfs.edit(form, [{"op": "fill_form", "fields": {"nope": "x"}}])


def test_word_and_powerpoint_and_csv(tmp_path):
    doc = readers.read(make_docx(tmp_path / "notes.docx"))
    assert doc["text"].splitlines()[:2] == ["# Release notes", "Fixed the login bug."]
    assert "| Area | Owner |" in doc["text"]
    deck = tmp_path / "deck.pptx"
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    with zipfile.ZipFile(deck, "w") as archive:
        for number, title in ((2, "Second"), (1, "First"), (10, "Tenth")):
            archive.writestr(f"ppt/slides/slide{number}.xml", f"<p:sld xmlns:p='x' {a}><a:p><a:r><a:t>{title}</a:t></a:r></a:p></p:sld>")
    slides = readers.read(deck)["text"]
    assert slides.index("First") < slides.index("Second") < slides.index("Tenth")
    table = tmp_path / "t.csv"
    sheets.edit(table, [{"op": "append_rows", "rows": [["name", "qty"], ["pens", 3]]}, {"op": "set", "cells": {"B2": 4}}], create=True)
    assert table.read_text().splitlines() == ["name,qty", "pens,4"]
    assert "2: pens | 4" in readers.read(table)["text"]


def test_workbook_edits_keep_formulas_and_formatting(tmp_path):
    book = tmp_path / "budget.xlsx"
    sheets.edit(
        book,
        [
            {"op": "append_rows", "rows": [["item", "cost"], ["rent", 1200], ["food", 300]]},
            {"op": "set", "cells": {"A4": "total", "B4": "=SUM(B2:B3)"}},
            {"op": "format", "range": "A1:B1", "bold": True, "fill": "#FFFF00", "width": 18},
            {"op": "add_sheet", "name": "Notes"},
            {"op": "freeze", "cell": "A2"},
        ],
        create=True,
    )
    from openpyxl import load_workbook

    ws = load_workbook(book)["Sheet"]
    assert ws["B4"].value == "=SUM(B2:B3)" and ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws.column_dimensions["A"].width == 18
    text = readers.read(book, sheet="Sheet")["text"]
    assert "B4==SUM(B2:B3)" in text and "A2=rent" in text
    assert readers.read(book, sheet="Sheet", cells="A2:B2")["text"].splitlines()[1] == "A2=rent  B2=1200"
    sheets.edit(book, [{"op": "delete_rows", "at": 3, "count": 1}, {"op": "rename_sheet", "sheet": "Notes", "name": "Read me"}])
    assert load_workbook(book).sheetnames == ["Sheet", "Read me"]


def test_a_workbook_with_charts_is_not_silently_damaged(tmp_path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    book = tmp_path / "chart.xlsx"
    wb = Workbook()
    ws = wb.active
    for row in ([1], [2], [3]):
        ws.append(row)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
    ws.add_chart(chart, "C1")
    wb.save(book)
    with pytest.raises(DocumentError, match="charts"):
        sheets.edit(book, [{"op": "set", "cells": {"A1": 9}}])
    sheets.edit(book, [{"op": "set", "cells": {"A1": 9}}], allow_lossy=True)


# ----- end to end, through the runtime ------------------------------------------


def _session(client, workspace, mode: str = "default") -> str:
    session_id = client.post("/api/sessions", json={"workspace": str(workspace)}).json()["id"]
    if mode != "default":
        client.patch(f"/api/sessions/{session_id}", json={"permission_mode": mode})
    return session_id


def _result(events, call_id):
    return next(e for e in of_type(events, "ToolResult") if e["payload"]["tool_call_id"] == call_id)["payload"]["return_value"]


@requires_runtime
def test_the_agent_reads_a_pdf_and_builds_a_workbook(client, workspace, script):
    make_pdf(workspace / "invoice.pdf", ["Invoice 42\nTotal due: 99.50"])
    script(
        [
            tool_call("c1", "read_document", {"path": "invoice.pdf"}),
            tool_call(
                "c2",
                "edit_spreadsheet",
                {"path": "ledger.xlsx", "create": True, "operations": [{"op": "append_rows", "rows": [["invoice", "amount"], [42, 99.5]]}, {"op": "set", "cells": {"B3": "=SUM(B2:B2)"}}]},
            ),
            text("Recorded invoice 42."),
        ]
    )
    session_id = _session(client, workspace)
    with client.websocket_connect(f"/api/sessions/{session_id}/ws") as ws:
        ws.receive_json()
        events = run_turn(ws, "log the invoice in a ledger")
    read = _result(events, "c1")
    assert not read["is_error"] and "Total due: 99.50" in read["output"]
    assert [a["payload"]["action"] for a in of_type(events, "ApprovalRequest")] == ["edit file"]
    assert not _result(events, "c2")["is_error"]
    from openpyxl import load_workbook

    assert load_workbook(workspace / "ledger.xlsx").active["B3"].value == "=SUM(B2:B2)"
    finished = of_type(events, "TurnFinished")[-1]["payload"]
    assert finished["files_touched"] == ["ledger.xlsx"] and finished["verified"] is True


@requires_runtime
def test_document_edits_respect_rejection_and_plan_mode(client, workspace, script):
    make_pdf(workspace / "a.pdf", ["keep me"])
    before = (workspace / "a.pdf").read_bytes()
    script([tool_call("c1", "edit_pdf", {"path": "a.pdf", "operations": [{"op": "add_text", "page": 1, "text": "x"}]}), text("ok")])
    session_id = _session(client, workspace)
    with client.websocket_connect(f"/api/sessions/{session_id}/ws") as ws:
        ws.receive_json()
        events = run_turn(ws, "stamp it", approve="reject")
    assert _result(events, "c1")["is_error"] and (workspace / "a.pdf").read_bytes() == before

    script([tool_call("c1", "edit_pdf", {"path": "a.pdf", "operations": [{"op": "rotate", "pages": "1"}]}), text("ok")])
    planner = _session(client, workspace, mode="plan")
    with client.websocket_connect(f"/api/sessions/{planner}/ws") as ws:
        ws.receive_json()
        events = run_turn(ws, "plan")
    result = _result(events, "c1")
    assert result["is_error"] and "not found" in json.dumps(result).lower()
    assert (workspace / "a.pdf").read_bytes() == before
